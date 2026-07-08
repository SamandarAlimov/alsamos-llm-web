"""office.py — jadval/hujjat eksporti (Excel .xlsx + .csv).

build_spreadsheet(): rows -> fayl. xlsx uchun openpyxl (lazy import),
csv uchun stdlib (bog'liqliksiz).
"""
from __future__ import annotations

import csv
import io
import os
import time
from typing import Dict, List, Optional, Sequence


def _safe_name(name: Optional[str], ext: str) -> str:
    base = (name or f"alsamos_{int(time.time())}").strip().replace(" ", "_")
    base = "".join(ch for ch in base if ch.isalnum() or ch in ("_", "-"))
    if not base:
        base = f"alsamos_{int(time.time())}"
    return f"{base}.{ext}"


def build_spreadsheet(rows: Sequence[Sequence], headers: Optional[List[str]] = None,
                      fmt: str = "xlsx", filename: Optional[str] = None,
                      out_dir: str = ".", sheet_name: str = "Alsamos") -> Dict:
    fmt = (fmt or "xlsx").lower()
    os.makedirs(out_dir, exist_ok=True)

    if fmt == "csv":
        name = _safe_name(filename, "csv")
        path = os.path.join(out_dir, name)
        buf = io.StringIO()
        w = csv.writer(buf)
        if headers:
            w.writerow(headers)
        for row in rows:
            w.writerow(list(row))
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(buf.getvalue())
        return {"path": path, "filename": name, "format": "csv", "rows": len(rows)}

    # xlsx
    from openpyxl import Workbook  # lazy
    from openpyxl.styles import Font, PatternFill
    name = _safe_name(filename, "xlsx")
    path = os.path.join(out_dir, name)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    r0 = 1
    if headers:
        ws.append(list(headers))
        fill = PatternFill("solid", fgColor="8B7CF6")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        r0 = 2
    for row in rows:
        ws.append(list(row))
    # ustun kengligini moslash
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 3, 60)
    ws.freeze_panes = f"A{r0}"
    wb.save(path)
    return {"path": path, "filename": name, "format": "xlsx", "rows": len(rows)}
